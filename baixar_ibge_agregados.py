"""Baixa agregados definitivos do Censo 2022 e recorta os setores de Goiânia."""
import argparse
import csv
import io
import json
from pathlib import Path
import re
import sys
import zipfile

from ibge_utils import AGREGADOS, MUNICIPIO, index, latest, local_records, now, obtain, save_json, sha256

STEMS={
    'basico':'Agregados_por_setores_basico_BR',
    'domicilio1':'Agregados_por_setores_caracteristicas_domicilio1_BR',
    'domicilio2':'Agregados_por_setores_caracteristicas_domicilio2_BR',
    'domicilio3':'Agregados_por_setores_caracteristicas_domicilio3_BR',
}


def discover():
    urls=index(AGREGADOS+'Agregados_por_Setor_csv/')
    items=[dict(id=key,tipo='agregados',url=latest(urls,stem,'.zip')) for key,stem in STEMS.items()]
    items.append(dict(id='dicionario_agregados',tipo='documentacao',url=latest(index(AGREGADOS),'dicionario_de_dados_agregados_por_setores_censitarios','.xlsx')))
    return items


def dictionary(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    rows=[]
    for ws in wb:
        if ws.title not in ['Dicionário Básico','Dicionário não PCT']:
            continue
        for values in ws.iter_rows(values_only=True):
            if len(values)==3:
                theme,code,description=values
                table='basico'
            else:
                _,theme,code,description=values
                match=re.search(r'Parte (\d)',str(theme))
                if not match:
                    continue
                table='domicilio'+match[1]
            if isinstance(code,str) and re.fullmatch(r'V\d+',code):
                rows.append(dict(tabela=table,variavel=code,descricao=str(description),tema=str(theme)))
    wb.close()
    if not rows:
        raise RuntimeError('Formato do dicionário não reconhecido.')
    return rows


def read_filtered(zip_path, municipality):
    import pandas as pd
    frames=[]; headers=None; total=0; members=[]
    with zipfile.ZipFile(zip_path) as z:
        csvs=[n for n in z.namelist() if n.lower().endswith('.csv') and not n.startswith('__MACOSX')]
        if not csvs:
            raise RuntimeError('ZIP não contém CSV.')
        for name in csvs:
            with z.open(name) as raw:
                sample=raw.read(65536)
            try:
                sample.decode('utf-8-sig'); encoding='utf-8-sig'
            except UnicodeDecodeError:
                encoding='latin-1'
            first=sample.decode(encoding).splitlines()[0]
            separator=';' if first.count(';')>first.count(',') else ','
            with z.open(name) as raw:
                for chunk in pd.read_csv(raw,sep=separator,encoding=encoding,dtype=str,keep_default_na=False,na_filter=False,chunksize=20000):
                    chunk.columns=[str(c).strip().upper() for c in chunk.columns]
                    chunk=chunk.loc[:,~chunk.columns.str.startswith('UNNAMED:')]
                    if 'SETOR' in chunk.columns and 'CD_SETOR' not in chunk.columns:
                        chunk=chunk.rename(columns={'SETOR':'CD_SETOR'})
                    if 'CD_SETOR' not in chunk.columns:
                        raise RuntimeError('CD_SETOR ausente em '+name)
                    if headers is None:
                        headers=list(chunk.columns)
                    if list(chunk.columns)!=headers:
                        raise RuntimeError('ZIP contém CSVs com esquemas diferentes.')
                    total+=len(chunk)
                    codes=chunk['CD_SETOR'].str.strip()
                    selected=chunk.loc[codes.str.startswith(municipality)].copy()
                    if len(selected):
                        selected['CD_SETOR']=selected['CD_SETOR'].str.strip()
                        frames.append(selected)
            members.append(name)
    if not frames:
        raise RuntimeError('Nenhum setor do município foi encontrado.')
    data=pd.concat(frames,ignore_index=True)
    if data['CD_SETOR'].duplicated().any() or not data['CD_SETOR'].str.fullmatch(municipality+r'\d{8}').all():
        raise RuntimeError('CD_SETOR duplicado ou fora do formato de 15 dígitos.')
    return data,dict(linhas_fonte=total,linhas_municipio=len(data),membros_csv=members)


def process(records,base):
    import pandas as pd
    source={r['id']:r for r in records}
    if not set(STEMS).union({'dicionario_agregados'}).issubset(source):
        raise RuntimeError('Manifesto sem todos os arquivos necessários.')
    output=Path(base)/'goiania'; output.mkdir(parents=True,exist_ok=True)
    definitions=dictionary(source['dicionario_agregados']['arquivo'])
    pd.DataFrame(definitions).to_csv(output/'dicionario_variaveis.csv',index=False,encoding='utf-8-sig')
    reports={}; tables={}
    for key in STEMS:
        print('  Filtrando '+key+'...',flush=True)
        data,report=read_filtered(source[key]['arquivo'],MUNICIPIO)
        known={r['variavel'] for r in definitions if r['tabela']==key}
        variables={c for c in data.columns if re.fullmatch(r'V\d+',c)}
        unknown=sorted(variables-known)
        if unknown:
            raise RuntimeError(f'Variáveis não documentadas em {key}: {unknown}')
        data['ANO_REFERENCIA']='2022'
        data['FONTE_ARQUIVO']=Path(source[key]['arquivo']).name
        path=output/f'{key}_setores_2022.csv'
        temp=path.with_suffix('.csv.part')
        data.to_csv(temp,index=False,encoding='utf-8-sig'); temp.replace(path)
        report.update(arquivo=str(path),sha256=sha256(path),fonte_sha256=source[key]['sha256'],
                      campos=list(data.columns),marcadores_nao_numericos={c:{k:int(v) for k,v in data[c].value_counts().loc[lambda s:~s.index.str.match(r'^-?\d+(?:[.,]\d+)?$')].items()} for c in sorted(variables) if (~data[c].str.match(r'^-?\d+(?:[.,]\d+)?$')).any()})
        reports[key]=report; tables[key]=data
        print(f"  {key}: {len(data)} setores",flush=True)
    selected_codes={'basico':[f'V{x:04}' for x in range(1,10)],'domicilio1':['V00001','V00004','V00005'],
                    'domicilio2':[f'V{x:05}' for x in [*range(111,119),199,200,201,*range(309,317),*range(397,403)]]}
    # Cada medida é acompanhada pela descrição e pela célula original. Não há imputação de x, hífen ou vazio.
    indicators=tables['basico'][['CD_SETOR']].copy()
    descriptions=[]
    for key,codes in selected_codes.items():
        missing=set(codes)-set(tables[key].columns)
        if missing:
            raise RuntimeError(f'Variáveis prioritárias ausentes em {key}: {missing}')
        prefixed=tables[key][['CD_SETOR',*codes]].rename(columns={c:key+'_'+c for c in codes})
        indicators=indicators.merge(prefixed,on='CD_SETOR',how='left',validate='one_to_one')
        for code in codes:
            definition=next(r for r in definitions if r['tabela']==key and r['variavel']==code)
            descriptions.append(dict(coluna=key+'_'+code,**definition))
    indicators['ANO_REFERENCIA']='2022'
    indicators.to_csv(output/'indicadores_prioritarios_setores_2022.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(descriptions).to_csv(output/'dicionario_indicadores_prioritarios.csv',index=False,encoding='utf-8-sig')
    basic=set(tables['basico']['CD_SETOR'])
    coverage={key:dict(sem_basico=sorted(set(df['CD_SETOR'])-basic),basico_sem_tema=sorted(basic-set(df['CD_SETOR']))) for key,df in tables.items()}
    save_json(Path(base)/'qualidade_agregados.json',dict(processado_em=now(),ano_referencia=2022,municipio=MUNICIPIO,tabelas=reports,cobertura=coverage,
               observacao='Células originais preservadas. Diferenças de cobertura não são convertidas em zero. Nenhuma agregação por bairro executada.'))
    return reports


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--saida',type=Path,default=Path('dados/ibge'))
    p.add_argument('--somente-baixar',action='store_true')
    p.add_argument('--processar-local',action='store_true')
    p.add_argument('--atualizar',action='store_true')
    args=p.parse_args()
    args.saida.mkdir(parents=True,exist_ok=True)
    records=local_records(args.saida,'manifesto_agregados.json') if args.processar_local else obtain(discover(),args.saida,'manifesto_agregados.json',args.atualizar)
    if not args.somente_baixar:
        process(records,args.saida)
    print('Agregados concluídos: '+str(args.saida.resolve()),flush=True)


if __name__=='__main__':
    try:
        main()
    except Exception as exc:
        print('ERRO: '+str(exc),file=sys.stderr)
        sys.exit(1)
