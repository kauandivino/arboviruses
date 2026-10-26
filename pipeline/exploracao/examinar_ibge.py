from pathlib import Path
from openpyxl import load_workbook

path=Path('dados/ibge/inspecao/dicionario_agregados.xlsx')
wb=load_workbook(path,read_only=True,data_only=True)
from collections import Counter
ws=wb['Dicionário não PCT']
counts=Counter()
for row in ws.iter_rows(values_only=True):
    counts[row[1]]+=1
    if row[2] in ['V00001','V00004','V00005'] or row[2] in [f'V{x:05}' for x in [*range(110,119),*range(200,215),*range(308,317),*range(397,403),*range(560,568)]]:
        print(row)
print('TEMAS',counts)
wb.close()
